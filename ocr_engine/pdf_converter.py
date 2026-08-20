import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pdf2docx import Converter
from .ocr import is_image_file, extract_text_lines_from_image, extract_text_lines_from_pdf

def create_docx_from_lines(lines_list, output_docx_path):
    """
    Creates a clean, formatted Word (.docx) document directly from extracted text lines.
    Ensures 100% of extracted text is preserved in docx without omission.
    """
    doc = Document()
    
    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    for page_lines in lines_list:
        if not page_lines:
            continue
            
        for item in page_lines:
            text = item['text'].strip()
            if not text:
                continue
                
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            # Check for bold headings / emphasis (e.g. Ingredients:, CONTAINS, MAY CONTAIN, All Caps)
            is_bold = (
                text.startswith("Ingredients:") or 
                text.startswith("CONTAINS") or 
                text.startswith("MAY CONTAIN") or 
                (text.isupper() and len(text) > 3)
            )
            
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            if is_bold:
                run.font.bold = True
                
    # Handle file locking on save
    target_docx_path = output_docx_path
    if os.path.exists(target_docx_path):
        try:
            with open(target_docx_path, 'a'):
                pass
        except PermissionError:
            base, ext = os.path.splitext(target_docx_path)
            target_docx_path = f"{base}_converted{ext}"
            print(f"[Warning] '{os.path.basename(output_docx_path)}' is currently open/locked. Saving to '{os.path.basename(target_docx_path)}' instead.")

    doc.save(target_docx_path)
    print(f"[Converter] Direct Word file created with {sum(len(p) for p in lines_list)} text line(s) saved to '{os.path.basename(target_docx_path)}'.")
    return True

def convert_image_or_scanned_pdf_to_docx(input_path, output_docx_path):
    """
    Directly converts an image file or scanned PDF file into a Word document using high-accuracy OCR line extraction.
    """
    print(f"[Converter] Extracting text and building Word document for '{os.path.basename(input_path)}'...")
    try:
        if is_image_file(input_path):
            lines = extract_text_lines_from_image(input_path)
            lines_list = [lines]
        else:
            lines_list = extract_text_lines_from_pdf(input_path)
            
        return create_docx_from_lines(lines_list, output_docx_path)
    except Exception as e:
        print(f"[Error] Direct image/scanned PDF to docx conversion failed: {e}")
        return False

def convert_pdf_to_docx(pdf_path, docx_path, is_scanned=False):
    """
    Converts a PDF file to a Word (.docx) file.
    If is_scanned is True, uses direct structured extraction to prevent text omission.
    For digital PDFs, uses pdf2docx layout converter.
    """
    if is_scanned:
        return convert_image_or_scanned_pdf_to_docx(pdf_path, docx_path)
        
    print(f"[Converter] Converting digital PDF '{os.path.basename(pdf_path)}' to Word format...")
    
    target_docx_path = docx_path
    if os.path.exists(target_docx_path):
        try:
            with open(target_docx_path, 'a'):
                pass
        except PermissionError:
            base, ext = os.path.splitext(target_docx_path)
            target_docx_path = f"{base}_converted{ext}"
            print(f"[Warning] '{os.path.basename(docx_path)}' is currently open/locked. Saving to '{os.path.basename(target_docx_path)}' instead.")

    try:
        cv = Converter(pdf_path)
        cv.convert(
            target_docx_path,
            start=0,
            end=None,
            ocr=0,
            parse_lattice_table=True,
            parse_stream_table=True,
            connected_border_tolerance=2.0
        )
        cv.close()
        print(f"[Converter] Digital PDF conversion complete. Saved to '{os.path.basename(target_docx_path)}'.")
        return True
    except Exception as e:
        print(f"[Error] Failed to convert PDF to Word: {e}")
        # Fall back to direct extraction
        return convert_image_or_scanned_pdf_to_docx(pdf_path, docx_path)
